"""
CV Screener - Document Reader

Reads text from:
    - TXT files
    - PDF files
    - DOCX files

Author: Daniel Orisakwe
Group Project: CV Screener
"""

from pathlib import Path
from pypdf import PdfReader
from docx import Document


def read_txt(file_path):
    """Read text from a TXT file."""
    try:
        with open(file_path, "r", encoding="utf-8") as file:
            return file.read()

    except UnicodeDecodeError:
        # Try another common encoding
        with open(file_path, "r", encoding="latin-1") as file:
            return file.read()


def read_pdf(file_path):
    """Read text from a PDF file."""
    text = []

    try:
        reader = PdfReader(file_path)

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    except Exception as error:
        raise RuntimeError(f"Could not read PDF: {error}")


def read_docx(file_path):
    """Read text from a DOCX file."""
    try:
        document = Document(file_path)

        paragraphs = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Also read text from tables
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())

                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return "\n".join(paragraphs)

    except Exception as error:
        raise RuntimeError(f"Could not read DOCX: {error}")


def read_document(file_path):
    """
    Automatically detect the file type and extract its text.

    Supported:
        .txt
        .pdf
        .docx
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if not path.is_file():
        raise ValueError(
            f"The path is not a file: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".txt":
        return read_txt(path)

    elif extension == ".pdf":
        return read_pdf(path)

    elif extension == ".docx":
        return read_docx(path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Only TXT, PDF and DOCX are supported."
        )


def clean_text(text):
    """Clean unnecessary spaces and blank lines."""
    lines = []

    for line in text.splitlines():
        line = " ".join(line.split())

        if line:
            lines.append(line)

    return "\n".join(lines)


if __name__ == "__main__":
    print("CV Document Reader")
    print("=" * 40)

    file_name = input("Enter CV file path: ").strip()

    try:
        text = read_document(file_name)
        text = clean_text(text)

        print("\nCV TEXT")
        print("=" * 40)
        print(text)

        print("\n" + "=" * 40)
        print(f"Characters extracted: {len(text)}")
        print(f"Words extracted: {len(text.split())}")

    except Exception as error:
        print(f"\nERROR: {error}")